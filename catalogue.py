"""Stage 0: load the NSCSTI 2.0 resource repository workbook into a clean catalogue.

The workbook is a metadata index, not the corpus. This script turns it into
data/catalogue.json, which later stages use to attach pillar / metric / stage
information to every document chunk and to answer "which documents are needed
for Metric N?" by direct lookup.

Run:  python catalogue.py
"""
import json
import re
from pathlib import Path

import openpyxl

WORKBOOK = Path("data/resource_repository.xlsx")
OUTPUT = Path("data/catalogue.json")

# Cautions written inside the template column of the workbook.
FLAG_PATTERNS = {
    "needs_revision": re.compile(r"not relevant|need[s]? to be revised", re.I),
    "make_a_copy": re.compile(r"needs? to be made a copy", re.I),
}


def clean(value):
    """Collapse whitespace and return None for empty cells."""
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def normalise_pillar(name):
    """'Pillar 1: A & B' and 'Pillar 1: A and B' must match."""
    name = clean(name)
    return name.replace("&", "and") if name else None


def pillar_number(name):
    match = re.match(r"Pillar\s+(\d+)", name or "")
    return int(match.group(1)) if match else None


def load_resource_mapping(ws):
    """Sheet 1: one row per (metric, supporting document)."""
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(cell is None for cell in row):
            continue
        pillar, metric_no, metric, doc, templates, stages = row[:6]
        if metric_no is None:
            continue
        template_lines = [t.strip() for t in str(templates or "").split("\n") if t.strip()]
        flags = sorted(
            flag for flag, pattern in FLAG_PATTERNS.items()
            if any(pattern.search(t) for t in template_lines)
        )
        records.append({
            "pillar_no": pillar_number(pillar),
            "pillar": normalise_pillar(pillar),
            "metric_no": int(metric_no),
            "metric": clean(metric),
            "supporting_document": clean(doc),
            "templates": template_lines,
            "stages": [s.strip() for s in str(stages or "").split(",") if s.strip()],
            "flags": flags,
        })
    return records


def load_checklist(ws):
    """Sheet 2: pillar, metric and evidence category appear only on the first
    row of each group, so they are forward-filled."""
    records = []
    pillar_no = pillar = metric_no = metric = category = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        p_no, p_name, m, cat, checklist = row[:5]
        if clean(checklist) is None:
            continue
        if p_name is not None:
            pillar = normalise_pillar(p_name)
            pillar_no = int(p_no) if p_no is not None else pillar_number(pillar)
        if m is not None:
            metric = clean(m)
            found = re.match(r"Metric\s+(\d+)", metric or "")
            metric_no = int(found.group(1)) if found else metric_no
        if cat is not None:
            category = clean(cat)
        records.append({
            "pillar_no": pillar_no,
            "pillar": pillar,
            "metric_no": metric_no,
            "metric": metric,
            "evidence_category": category,
            "checklist_item": clean(checklist),
        })
    return records


def main():
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    mapping = load_resource_mapping(wb.worksheets[0])
    checklist = load_checklist(wb.worksheets[1])
    OUTPUT.write_text(
        json.dumps({"resource_mapping": mapping, "checklist": checklist},
                   indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    print(f"resource_mapping rows : {len(mapping)}")
    print(f"checklist rows        : {len(checklist)}")
    print(f"metrics (mapping)     : {len({r['metric_no'] for r in mapping})}")
    print(f"metrics (checklist)   : {len({r['metric_no'] for r in checklist})}")
    print(f"flagged rows          : {sum(1 for r in mapping if r['flags'])}")
    print(f"written to {OUTPUT}")


if __name__ == "__main__":
    main()
